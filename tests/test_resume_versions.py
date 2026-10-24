import asyncio
import io
from datetime import UTC, datetime
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest

from app.domain.resume_versions import approved_version_rules, draft_version_rules, version_metadata
from app.candidates.resume_export import export_resume
from app.candidates.resume_edit import content_digest, render_resume_text_diff, validate_resume_edit
from app.database.models import Job, ResumeProfile
from app.matching.engine import choose_resume
from app.telegram import handlers


def test_legacy_resume_is_compatible_and_new_version_requires_approval():
    assert version_metadata({}) == {"status": "approved", "legacy": True}
    rules = draft_version_rules({"format": "ats"}, parent_id=7)
    assert rules == {"format": "ats", "cv_version_v1": {"status": "draft", "parent_id": 7}}
    approved = approved_version_rules(rules, now=datetime(2026, 10, 23, tzinfo=UTC))
    assert approved["format"] == "ats"
    assert approved["cv_version_v1"] == {
        "status": "approved", "parent_id": 7, "approved_at": "2026-10-23T00:00:00Z",
    }


def test_draft_resume_is_never_selected_for_matching_or_application_preparation():
    job = Job(title="SRE", company="Example", description="Kubernetes", country="US")
    approved = ResumeProfile(id=1, profile_id=1, name="Base", role="DevOps", language="en", rules={})
    draft = ResumeProfile(
        id=2, profile_id=1, name="Tailored", role="SRE", language="en",
        content="Kubernetes", rules=draft_version_rules({}, parent_id=1),
    )
    assert choose_resume(job, [approved, draft]) == 1


def test_resume_list_discloses_draft_and_approved_state(monkeypatch):
    rows = [
        SimpleNamespace(id=1, name="Legacy", role="DevOps", language="en", source="imported", rules={}),
        SimpleNamespace(
            id=2, name="SRE draft", role="SRE", language="en", source="version",
            rules=draft_version_rules({}, parent_id=1),
        ),
    ]
    monkeypatch.setattr(handlers, "_require_user", AsyncMock(return_value=SimpleNamespace(id=5)))
    monkeypatch.setattr(handlers.repo, "active_profile", AsyncMock(return_value=SimpleNamespace(id=7)))
    monkeypatch.setattr(handlers.repo, "list_resumes", AsyncMock(return_value=rows))
    message = SimpleNamespace(answer=AsyncMock())

    asyncio.run(handlers.resumes(message))

    rendered = message.answer.await_args.args[0]
    assert "Legacy — DevOps — en — imported — ✅ approved" in rendered
    assert "SRE draft — SRE — en — version — 📝 draft" in rendered


def test_resume_duplicate_and_approve_are_owner_profile_scoped(monkeypatch):
    profile = SimpleNamespace(id=7)
    monkeypatch.setattr(handlers, "_require_user", AsyncMock(return_value=SimpleNamespace(id=5)))
    monkeypatch.setattr(handlers.repo, "active_profile", AsyncMock(return_value=profile))
    duplicate = AsyncMock(return_value=SimpleNamespace(id=12))
    approve = AsyncMock(return_value=True)
    monkeypatch.setattr(handlers.repo, "duplicate_resume_version", duplicate)
    monkeypatch.setattr(handlers.repo, "approve_resume_version", approve)

    duplicate_message = SimpleNamespace(text="/resumeduplicate 3 SRE EN", answer=AsyncMock())
    asyncio.run(handlers.resume_duplicate_command(duplicate_message))
    duplicate.assert_awaited_once_with(7, 3, "SRE EN")
    assert "/resumeapprove 12" in duplicate_message.answer.await_args.args[0]

    approve_message = SimpleNamespace(text="/resumeapprove 12", answer=AsyncMock())
    asyncio.run(handlers.resume_approve_command(approve_message))
    approve.assert_awaited_once_with(7, 12)
    assert "утверждена" in approve_message.answer.await_args.args[0]


def test_approved_resume_exports_exact_text_to_txt_and_ats_friendly_docx():
    from docx import Document

    resume = SimpleNamespace(
        id=4, name="SRE / EN", role="Site Reliability Engineer", language="en",
        content="Summary\nLinux & Kubernetes", rules=approved_version_rules({}, now=datetime(2026, 10, 23, tzinfo=UTC)),
    )
    txt_name, txt = export_resume(resume, "txt")
    assert txt_name == "SRE-EN.txt"
    assert txt.decode() == resume.content

    docx_name, docx = export_resume(resume, "docx")
    assert docx_name == "SRE-EN-ats.docx"
    rendered = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(docx)).paragraphs)
    assert rendered == "SRE / EN\nSite Reliability Engineer · EN\nSummary\nLinux & Kubernetes"


def test_modern_docx_and_pdf_preserve_ru_en_text_without_new_claims():
    from docx import Document
    from pypdf import PdfReader

    resume = SimpleNamespace(
        id=5, name="SRE RU", role="Инженер надёжности", language="ru",
        content="Опыт\nLinux и Kubernetes", rules={},
    )
    docx_name, docx = export_resume(resume, "docx", "modern")
    assert docx_name == "SRE-RU-modern.docx"
    doc = Document(io.BytesIO(docx))
    assert "Опыт\nLinux и Kubernetes" in "\n".join(item.text for item in doc.paragraphs)
    assert str(doc.paragraphs[0].runs[0].font.color.rgb) == "1F4E79"

    pdf_name, pdf = export_resume(resume, "pdf", "modern")
    assert pdf_name == "SRE-RU-modern.pdf"
    assert pdf.startswith(b"%PDF-")
    extracted = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)
    assert "Инженер надёжности" in extracted
    assert "Linux и Kubernetes" in extracted


def test_draft_or_empty_resume_cannot_be_exported():
    draft = SimpleNamespace(name="Draft", content="Private", rules=draft_version_rules({}, parent_id=1))
    with pytest.raises(ValueError, match="утверди"):
        export_resume(draft, "txt")
    empty = SimpleNamespace(name="Empty", content="", rules={})
    with pytest.raises(ValueError, match="нет текста"):
        export_resume(empty, "docx")


def test_resume_export_command_uses_owner_scoped_lookup(monkeypatch):
    resume = SimpleNamespace(id=4, name="SRE", role="SRE", language="en", content="Linux", rules={})
    monkeypatch.setattr(handlers, "_require_user", AsyncMock(return_value=SimpleNamespace(id=5)))
    monkeypatch.setattr(handlers.repo, "active_profile", AsyncMock(return_value=SimpleNamespace(id=7)))
    lookup = AsyncMock(return_value=resume)
    monkeypatch.setattr(handlers.repo, "get_profile_resume", lookup)
    message = SimpleNamespace(
        text="/resumeexport 4 txt", answer=AsyncMock(), answer_document=AsyncMock(),
    )

    asyncio.run(handlers.resume_export_command(message))

    lookup.assert_awaited_once_with(7, 4)
    message.answer.assert_not_awaited()
    document = message.answer_document.await_args.args[0]
    assert document.filename == "SRE.txt"


def test_resume_edit_validation_and_diff_are_grounded_in_exact_text():
    proposed = validate_resume_edit("Linux\nDocker", "Linux\nKubernetes")
    diff = render_resume_text_diff("Linux\nDocker", proposed)
    assert "-Docker" in diff
    assert "+Kubernetes" in diff
    assert content_digest("Linux") != content_digest("linux")
    with pytest.raises(ValueError, match="совпадает"):
        validate_resume_edit("Linux", " Linux ")


def test_resume_edit_preview_does_not_write_and_confirmation_creates_draft(monkeypatch):
    handlers._PENDING_RESUME_EDITS.clear()
    resume = SimpleNamespace(id=4, content="Linux\nDocker")
    monkeypatch.setattr(handlers, "_require_user", AsyncMock(return_value=SimpleNamespace(id=5)))
    monkeypatch.setattr(handlers.repo, "active_profile", AsyncMock(return_value=SimpleNamespace(id=7)))
    monkeypatch.setattr(handlers.repo, "get_profile_resume", AsyncMock(return_value=resume))
    create = AsyncMock(return_value=SimpleNamespace(id=9))
    monkeypatch.setattr(handlers.repo, "create_resume_revision", create)
    preview_message = SimpleNamespace(
        text="/resumeedit 4 Linux\nKubernetes", answer=AsyncMock(),
    )

    asyncio.run(handlers.resume_edit_command(preview_message))

    create.assert_not_awaited()
    rendered = preview_message.answer.await_args.args[0]
    assert "-Docker" in rendered and "+Kubernetes" in rendered
    token = next(iter(handlers._PENDING_RESUME_EDITS))
    confirm = SimpleNamespace(text=f"/resumeeditconfirm {token}", answer=AsyncMock())
    asyncio.run(handlers.resume_edit_confirm_command(confirm))
    create.assert_awaited_once_with(
        7, 4, expected_content_digest=content_digest("Linux\nDocker"), content="Linux\nKubernetes",
    )
    assert "/resumeapprove 9" in confirm.answer.await_args.args[0]

    replay = SimpleNamespace(text=f"/resumeeditconfirm {token}", answer=AsyncMock())
    asyncio.run(handlers.resume_edit_confirm_command(replay))
    assert "не найден или истёк" in replay.answer.await_args.args[0]
