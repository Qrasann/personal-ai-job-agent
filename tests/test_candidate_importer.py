from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from docx import Document

from app.candidates.importer import (
    MAX_DOCUMENT_BYTES,
    FactProposal,
    ResumeImportError,
    canonical_fact,
    extract_resume_text,
    guess_language,
    propose_facts,
    resume_fingerprint,
    review_proposals,
)


def test_text_import_supports_utf8_and_windows_1251():
    assert "Навыки" in extract_resume_text("resume.txt", "Навыки\nLinux".encode())
    encoded = "Опыт работы\nLinux".encode("cp1251")
    assert "Опыт работы" in extract_resume_text("resume.txt", encoded)


def test_text_import_rejects_empty_unsupported_and_oversized_files():
    with pytest.raises(ResumeImportError, match="пуст"):
        extract_resume_text("resume.txt", b"  \n")
    with pytest.raises(ResumeImportError, match="только PDF"):
        extract_resume_text("resume.rtf", b"text")
    with pytest.raises(ResumeImportError, match="8 МБ"):
        extract_resume_text("resume.txt", b"x" * (MAX_DOCUMENT_BYTES + 1))


def test_docx_import_reads_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Skills")
    doc.add_paragraph("Linux, Docker")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Experience"
    table.cell(0, 1).text = "Managed Kubernetes"
    output = io.BytesIO()
    doc.save(output)
    text = extract_resume_text("profile.docx", output.getvalue())
    assert "Linux, Docker" in text
    assert "Experience | Managed Kubernetes" in text


def test_pdf_import_extracts_text_without_ocr(monkeypatch):
    class FakePage:
        def extract_text(self):
            return "Skills\nLinux"

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

        def __init__(self, stream, strict):
            assert stream.read(5) == b"%PDF-"
            assert strict is False

    monkeypatch.setattr("pypdf.PdfReader", FakeReader)
    assert extract_resume_text("profile.pdf", b"%PDF-fake") == "Skills\nLinux"


def test_proposals_are_section_scoped_and_exclude_contact_data():
    text = """Alex Example
alex@example.com
Skills
- Linux, Docker, PostgreSQL
- Linux, Docker, PostgreSQL
Experience
- Operated Kubernetes clusters for production workloads
Phone: +1 555 123 4567
Education
- Bachelor of Computer Science
"""
    proposals = propose_facts(text)
    assert [item.category for item in proposals] == ["skills", "experience", "education"]
    assert proposals[0].confidence == 0.92
    assert all("@" not in item.value and "555" not in item.value for item in proposals)
    assert all(item.experience_type == "unknown" for item in proposals)


def test_proposal_review_skips_exact_duplicates_and_flags_overlaps_without_overwriting():
    proposals = (
        FactProposal("skills", "Linux and Docker", 0.92, "skills"),
        FactProposal("skills", "PostgreSQL", 0.92, "skills"),
    )
    existing = [
        SimpleNamespace(id=3, value="PostgreSQL"),
        SimpleNamespace(id=4, value="Linux Docker in production"),
    ]
    review = review_proposals(proposals, existing)
    assert len(review.duplicates) == 1
    assert len(review.new) == 1
    assert [(item.existing_id, item.proposal.value) for item in review.overlaps] == [(4, "Linux and Docker")]


def test_candidate_import_helpers_are_stable():
    assert canonical_fact(" Linux, Docker! ") == "linux docker"
    assert guess_language("Опыт работы Linux") == "ru"
    assert guess_language("Linux and Docker experience") == "en"
    assert resume_fingerprint(" Linux   Docker ") == resume_fingerprint("linux docker")
