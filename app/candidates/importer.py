from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath


MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_TEXT_CHARS = 80_000
MAX_PDF_PAGES = 30
MAX_DOCX_ENTRIES = 500
MAX_DOCX_UNPACKED_BYTES = 24 * 1024 * 1024
MAX_FACT_PROPOSALS = 30

_HEADINGS = {
    "skills": "skills", "technical skills": "skills", "key skills": "skills",
    "technologies": "skills", "technology stack": "skills", "core competencies": "skills",
    "навыки": "skills", "технические навыки": "skills", "технологии": "skills",
    "компетенции": "skills", "стек технологий": "skills",
    "experience": "experience", "work experience": "experience",
    "professional experience": "experience", "employment history": "experience",
    "опыт": "experience", "опыт работы": "experience", "профессиональный опыт": "experience",
    "projects": "experience", "personal projects": "experience", "проекты": "experience",
    "education": "education", "certifications": "education", "certificates": "education",
    "образование": "education", "сертификаты": "education", "сертификация": "education",
}
_PRIVATE_LINE = re.compile(
    r"(?:e-?mail|телефон|phone|паспорт|passport|дата рождения|date of birth|"
    r"address|адрес|снилс|инн|social security|tax id)", re.I,
)
_EMAIL = re.compile(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b")
_PHONE = re.compile(r"(?:\+?\d[\d ().-]{7,}\d)")
_WORD = re.compile(r"[\w+#.-]+", re.UNICODE)


class ResumeImportError(ValueError):
    """A safe, user-displayable reason that a resume could not be imported."""


@dataclass(frozen=True)
class FactProposal:
    category: str
    value: str
    confidence: float
    section: str
    experience_type: str = "unknown"


@dataclass(frozen=True)
class FactOverlap:
    proposal: FactProposal
    existing_id: int
    existing_value: str


@dataclass(frozen=True)
class ProposalReview:
    new: tuple[FactProposal, ...]
    duplicates: tuple[FactProposal, ...]
    overlaps: tuple[FactOverlap, ...]


def _bounded_text(text: str) -> str:
    cleaned = "\n".join(line.strip() for line in text.replace("\x00", "").replace("\r", "\n").splitlines())
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if not cleaned:
        raise ResumeImportError("Файл пуст или не содержит извлекаемого текста. Возможно, PDF состоит из сканов без текстового слоя.")
    if len(cleaned) > MAX_TEXT_CHARS:
        raise ResumeImportError(f"Документ слишком большой по тексту (лимит {MAX_TEXT_CHARS} символов).")
    return cleaned


def extract_resume_text(filename: str, payload: bytes) -> str:
    if not payload:
        raise ResumeImportError("Файл пуст.")
    if len(payload) > MAX_DOCUMENT_BYTES:
        raise ResumeImportError("Размер файла превышает лимит 8 МБ.")
    suffix = PurePath(filename or "").suffix.casefold()

    if suffix in {".txt", ".text"}:
        if b"\x00" in payload:
            raise ResumeImportError("TXT содержит бинарные данные.")
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = payload.decode("cp1251")
            except UnicodeDecodeError as exc:
                raise ResumeImportError("TXT должен быть в UTF-8 или Windows-1251.") from exc
        return _bounded_text(text)

    if suffix == ".pdf":
        if not payload.startswith(b"%PDF-"):
            raise ResumeImportError("Файл с расширением PDF не имеет сигнатуры PDF.")
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(payload), strict=False)
            if reader.is_encrypted:
                raise ResumeImportError("Зашифрованный PDF не поддерживается.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ResumeImportError(f"PDF содержит больше {MAX_PDF_PAGES} страниц.")
            parts = []
            total = 0
            for page in reader.pages:
                part = page.extract_text() or ""
                total += len(part)
                if total > MAX_TEXT_CHARS:
                    raise ResumeImportError(f"Документ слишком большой по тексту (лимит {MAX_TEXT_CHARS} символов).")
                parts.append(part)
        except ResumeImportError:
            raise
        except Exception as exc:
            raise ResumeImportError("Не удалось прочитать PDF. Проверьте файл и пароль.") from exc
        return _bounded_text("\n".join(parts))

    if suffix == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_DOCX_ENTRIES:
                    raise ResumeImportError("DOCX содержит слишком много частей.")
                if sum(item.file_size for item in entries) > MAX_DOCX_UNPACKED_BYTES:
                    raise ResumeImportError("Внутренний размер DOCX превышает защитный лимит.")
                if "word/document.xml" not in archive.namelist():
                    raise ResumeImportError("Файл не является документом DOCX.")
            from docx import Document

            document = Document(io.BytesIO(payload))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return _bounded_text("\n".join(parts))
        except ResumeImportError:
            raise
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise ResumeImportError("Не удалось прочитать DOCX. Проверьте файл.") from exc
        except Exception as exc:
            raise ResumeImportError("Не удалось прочитать DOCX. Проверьте файл.") from exc

    raise ResumeImportError("Поддерживаются только PDF, DOCX и TXT.")


def propose_facts(text: str) -> tuple[FactProposal, ...]:
    proposals: list[FactProposal] = []
    section: str | None = None
    seen: set[str] = set()
    for raw_line in text.splitlines():
        line = re.sub(r"^\s*(?:[•*▪‣–—-]+|\d+[.)])\s*", "", raw_line).strip()
        heading = re.sub(r"[\s:：]+$", "", line).casefold()
        if heading in _HEADINGS:
            section = _HEADINGS[heading]
            continue
        if line.startswith("#"):
            heading = line.lstrip("# ").casefold().rstrip(":")
            section = _HEADINGS.get(heading)
            continue
        if section is None or len(line) < 4 or len(line) > 300:
            continue
        if _PRIVATE_LINE.search(line) or _EMAIL.search(line) or _PHONE.search(line):
            continue
        canonical = canonical_fact(line)
        if not canonical or canonical in seen:
            continue
        seen.add(canonical)
        confidence = {"skills": 0.92, "experience": 0.82, "education": 0.90}[section]
        proposals.append(FactProposal(section, line, confidence, section))
        if len(proposals) >= MAX_FACT_PROPOSALS:
            break
    return tuple(proposals)


def canonical_fact(value: str) -> str:
    return " ".join(_WORD.findall((value or "").casefold()))


def _fact_tokens(value: str) -> set[str]:
    stop_words = {"and", "or", "the", "in", "with", "for", "of", "на", "и", "в", "с", "по"}
    return {token for token in _WORD.findall((value or "").casefold()) if len(token) > 1 and token not in stop_words}


def review_proposals(proposals, existing_facts) -> ProposalReview:
    existing = [(fact, canonical_fact(fact.value)) for fact in existing_facts]
    new: list[FactProposal] = []
    duplicates: list[FactProposal] = []
    overlaps: list[FactOverlap] = []
    seen_new: set[str] = set()
    for proposal in proposals:
        canonical = canonical_fact(proposal.value)
        if not canonical or canonical in seen_new:
            continue
        seen_new.add(canonical)
        exact = next((fact for fact, value in existing if value == canonical), None)
        if exact:
            duplicates.append(proposal)
            continue
        proposal_tokens = _fact_tokens(proposal.value)
        for fact, _ in existing:
            existing_tokens = _fact_tokens(fact.value)
            smaller, larger = sorted((proposal_tokens, existing_tokens), key=len)
            if len(smaller) >= 2 and len(smaller & larger) / len(smaller) >= 0.8:
                overlaps.append(FactOverlap(proposal, int(fact.id), str(fact.value)))
                break
        new.append(proposal)
    return ProposalReview(tuple(new), tuple(duplicates), tuple(overlaps))


def guess_language(text: str) -> str:
    letters = [char for char in text if char.isalpha()]
    if not letters:
        return "und"
    cyrillic = sum("\u0400" <= char <= "\u04ff" for char in letters)
    return "ru" if cyrillic / len(letters) >= 0.2 else "en"


def resume_fingerprint(text: str) -> str:
    normalized = " ".join((text or "").casefold().split())
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()
