from __future__ import annotations

import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.domain.resume_versions import version_metadata


def _safe_filename(value: str) -> str:
    stem = re.sub(r"[^\w.-]+", "-", value.strip(), flags=re.UNICODE).strip("-.")
    return stem[:80] or "resume"


def _pdf_font() -> str:
    name = "JobAgentUnicode"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    candidates = (
        Path("/usr/share/fonts/noto/NotoSans-Regular.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    font_path = next((path for path in candidates if path.is_file()), None)
    if font_path is None:
        raise ValueError("PDF export требует локальный Unicode-шрифт Noto Sans или DejaVu Sans.")
    pdfmetrics.registerFont(TTFont(name, str(font_path)))
    return name


def export_resume(resume, file_format: str, style: str = "ats") -> tuple[str, bytes]:
    """Export an explicitly approved stored CV without adding inferred claims."""
    file_format = file_format.casefold()
    style = style.casefold()
    if file_format not in {"txt", "docx", "pdf"}:
        raise ValueError("Формат экспорта: txt, docx или pdf.")
    if style not in {"ats", "modern"}:
        raise ValueError("Шаблон экспорта: ats или modern.")
    if version_metadata(getattr(resume, "rules", None))["status"] != "approved":
        raise ValueError("Сначала утверди эту версию CV через /resumeapprove.")
    content = str(getattr(resume, "content", "") or "").strip()
    if not content:
        raise ValueError("В этой версии CV нет текста для экспорта.")
    filename = _safe_filename(str(getattr(resume, "name", "") or "resume"))
    if file_format == "txt":
        return f"{filename}.txt", content.encode("utf-8")

    name = str(getattr(resume, "name", "") or "Resume")
    metadata = " · ".join(filter(None, (
        str(getattr(resume, "role", "") or "").strip(),
        str(getattr(resume, "language", "") or "").upper().strip(),
    )))
    if file_format == "docx":
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.65 if style == "modern" else 0.8)
        heading = document.add_heading(name, level=1)
        if style == "modern":
            heading.runs[0].font.color.rgb = RGBColor(31, 78, 121)
            heading.runs[0].font.size = Pt(24)
        if metadata:
            paragraph = document.add_paragraph(metadata)
            if style == "modern":
                paragraph.runs[0].font.color.rgb = RGBColor(89, 89, 89)
        for line in content.splitlines():
            document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        return f"{filename}-{style}.docx", output.getvalue()

    output = io.BytesIO()
    font = _pdf_font()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title=name, author="Personal AI Job Agent",
    )
    styles = getSampleStyleSheet()
    accent = HexColor("#1F4E79") if style == "modern" else HexColor("#111111")
    title_style = ParagraphStyle(
        "CVTitle", parent=styles["Title"], fontName=font,
        fontSize=22 if style == "modern" else 18, leading=26, textColor=accent, alignment=TA_LEFT,
    )
    meta_style = ParagraphStyle(
        "CVMeta", parent=styles["Normal"], fontName=font,
        fontSize=10, leading=14, textColor=HexColor("#555555"),
    )
    body_style = ParagraphStyle(
        "CVBody", parent=styles["Normal"], fontName=font,
        fontSize=10.5, leading=15, spaceAfter=4,
    )
    story = [Paragraph(escape(name), title_style)]
    if metadata:
        story.extend([Paragraph(escape(metadata), meta_style), Spacer(1, 8)])
    for line in content.splitlines():
        story.append(Paragraph(escape(line) or "&#160;", body_style))
    document.build(story)
    return f"{filename}-{style}.pdf", output.getvalue()
